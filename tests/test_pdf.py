"""Tests for pdf.py (PyMuPDF is a core dependency; nothing is skipped)."""

import os

import fitz
import pytest
from pptx import Presentation

from appt_ocr.pdf import convert_pdf_to_pptx
from tests.conftest import picture_shapes


def test_one_slide_per_page_at_the_page_size(sample_pdf_path):
    out = convert_pdf_to_pptx(sample_pdf_path, dpi=72)
    try:
        prs = Presentation(out)
        assert len(prs.slides) == 1
        assert (prs.slide_width, prs.slide_height) == (
            int(200 * 914400 / 72),
            int(150 * 914400 / 72),
        )
        (pic,) = picture_shapes(prs.slides[0])
        assert (pic.left, pic.top, pic.width, pic.height) == (
            0,
            0,
            prs.slide_width,
            prs.slide_height,
        )
        assert pic.image.size == (200, 150)  # rendered at 72 dpi: one pixel per point
    finally:
        os.unlink(out)


def test_dpi_scales_the_rendering(sample_pdf_path):
    out = convert_pdf_to_pptx(sample_pdf_path, dpi=144)
    try:
        (pic,) = picture_shapes(Presentation(out).slides[0])
        assert pic.image.size == (400, 300)
    finally:
        os.unlink(out)


def test_multi_page(temp_dir):
    doc = fitz.open()
    for _ in range(3):
        doc.new_page(width=100, height=100)
    path = os.path.join(temp_dir, "three.pdf")
    doc.save(path)
    doc.close()
    out = convert_pdf_to_pptx(path, dpi=36)
    try:
        assert len(Presentation(out).slides) == 3
    finally:
        os.unlink(out)


def test_an_empty_pdf_is_refused(temp_dir):
    # PyMuPDF will not save a zero-page document, so write one by hand; MuPDF
    # repairs the missing xref on open. Until 3.1.0 this was an IndexError.
    path = os.path.join(temp_dir, "empty.pdf")
    with open(path, "wb") as fh:
        fh.write(
            b"%PDF-1.4\n"
            b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
            b"2 0 obj << /Type /Pages /Kids [] /Count 0 >> endobj\n"
            b"trailer << /Root 1 0 R >>\n%%EOF\n"
        )
    with pytest.raises(ValueError, match="no pages"):
        convert_pdf_to_pptx(path)


def test_a_missing_file_raises(temp_dir):
    with pytest.raises(
        (FileNotFoundError, RuntimeError)
    ):  # pymupdf's own FileNotFoundError
        convert_pdf_to_pptx(os.path.join(temp_dir, "missing.pdf"))
